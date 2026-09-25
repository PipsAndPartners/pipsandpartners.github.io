import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import discovery.shortlist as sl
from outreach.contacted import SHEET_COLS, norm_url, record_contacted, tick_sheet
from scripts.backfill_contacted import sent_records

URL = "https://www.youtube.com/channel/UCnH3mGTw4w7rjOYIfGAVV7Q"


def _xlsx(path: Path, header: list, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def _sheet(path: Path) -> tuple[list, list[tuple]]:
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    return list(rows[0]), rows[1:]


class NormUrlTests(unittest.TestCase):
    def test_scheme_www_slash_case(self):
        self.assertEqual(norm_url("HTTPS://www.YouTube.com/channel/UCx/"), "youtube.com/channel/ucx")
        self.assertEqual(norm_url(None), "")


class TickSheetTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = _xlsx(Path(self.dir.name) / "s.xlsx", ["email", "source_url", "follow up note"], [
            ["a@x.com", "", "gọi lại thứ 2"],
            ["mailto:A@x.com", "", None],
            ["b@x.com", URL + "/", None],
            ["c@x.com", "", None],
        ])

    def tearDown(self):
        self.dir.cleanup()

    def test_ticks_by_email_or_url_and_keeps_manual_edits(self):
        state = {"a@x.com": {"contacted": True, "follow_up_at": date(2026, 10, 1), "note": "replied"},
                 norm_url(URL): {"contacted": True, "website": "azfin.vn"}}
        self.assertEqual(tick_sheet(self.path, state), 3)
        header, rows = _sheet(self.path)
        self.assertEqual(header, ["email", "source_url", "follow up note", "website",
                                  "already contact", "follow up"])
        col = {h: header.index(h) for h in header}
        self.assertEqual([r[col["already contact"]] for r in rows], ["Yes", "Yes", "Yes", None])
        self.assertEqual(rows[0][col["follow up note"]], "gọi lại thứ 2")  # not overwritten
        self.assertEqual(rows[1][col["follow up note"]], "replied")
        self.assertEqual(rows[2][col["website"]], "azfin.vn")

    def test_rerun_adds_no_duplicate_columns_or_dropdowns(self):
        tick_sheet(self.path, {})
        tick_sheet(self.path, {})
        ws = load_workbook(self.path).active
        header = [c.value for c in ws[1]]
        self.assertEqual(len(header), len(set(header)))
        self.assertTrue(set(SHEET_COLS) <= set(header))
        self.assertEqual(len(ws.data_validations.dataValidation), 1)


class RecordContactedTests(unittest.TestCase):
    def test_sheet_still_ticked_when_db_is_down(self):
        with tempfile.TemporaryDirectory() as d:
            path = _xlsx(Path(d) / "m.xlsx", ["email", "Content for mail"], [["a@x.com", "hi"], ["b@x.com", "hi"]])
            with patch("db.connection.get_conn", side_effect=RuntimeError("no db")):
                record_contacted([{"to": "a@x.com", "source_url": None}], "t", path)
            header, rows = _sheet(path)
            c = header.index("already contact")
            self.assertEqual([r[c] for r in rows], ["Yes", None])
            self.assertIsNotNone(rows[0][header.index("follow up")])


class SentRecordsTests(unittest.TestCase):
    def test_parses_sent_items_export(self):
        with tempfile.TemporaryDirectory() as d:
            path = _xlsx(Path(d) / "sent.xlsx",
                         ["Time", "Name", "Phone", "mail", "facebook", "tiktok", "youtube", "zalo/telegram",
                          "Client reply"],
                         [["2026-09-21 11:52", "AzFin", None, "'Academy@azfin.vn'", None, None, URL, None, "yes"],
                          ["2026-09-21 11:53", "blank", None, None, None, None, None, None, None]])
            recs = sent_records(path)
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0]["to"], "academy@azfin.vn")
        self.assertEqual(recs[0]["urls"], [URL])
        self.assertEqual(recs[0]["note"], "replied")
        self.assertEqual(recs[0]["at"].date(), date(2026, 9, 21))


class ShortlistColsTests(unittest.TestCase):
    def test_row_values_align_with_cols(self):
        r = {"platform": "youtube", "display_name": "AzFin", "handle": "@azfin", "follower_count": 1,
             "creator_relevance": 1, "brand_safety": 1, "engagement_rate": 1, "composite_score": 1,
             "email_norm": "a@x.com", "phone_norm": None, "contacts": None, "source_url": URL,
             "website": "azfin.vn", "contacted_at": "2026-09-21", "follow_up_at": date(2026, 9, 28),
             "follow_up_note": None}
        row = dict(zip(sl.COLS, sl._row_values(1, r, None)))
        self.assertEqual(len(sl._row_values(1, r, None)), len(sl.COLS))
        self.assertEqual(row["website"], "azfin.vn")
        self.assertEqual(row["already contact"], "Yes")
        self.assertEqual(row["source_url"], URL)


if __name__ == "__main__":
    unittest.main()
