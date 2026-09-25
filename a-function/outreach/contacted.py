"""Already-contacted tracking: who got an outreach email, and when to follow up.

Source of truth is `entities.contacted_at` / `contacted_batch` / `follow_up_at` /
`follow_up_note` (db/12_outreach_status.sql). outreach.createmail calls
record_contacted() after writing its .eml files, which marks the DB and ticks
"already contact" = Yes in the sheet it read. scripts/backfill_contacted replays
older batches + the Outlook Sent Items export through the same mark_contacted().

    # add website / already contact / follow up columns to any sheet, from the DB:
    python -m outreach.contacted results/ib_master.xlsx results/ib_master_email.xlsx
    # follow-up worklist of everyone contacted, soonest follow-up first:
    python -m outreach.contacted --tracker results/outreach_tracker.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from discovery.shortlist import candidate_ids

FOLLOW_UP_DAYS = 7
SHEET_COLS = ("website", "already contact", "follow up", "follow up note")
_NOT_IDS = {"profile.php", "people", "pages", "groups"}  # URL path words, not account ids

MARK_SQL = """
    UPDATE entities SET
        contacted_at    = coalesce(contacted_at, %(at)s),
        contacted_batch = coalesce(contacted_batch, %(batch)s),
        email_norm      = coalesce(email_norm, %(email)s),
        follow_up_at    = coalesce(follow_up_at, %(follow)s),
        follow_up_note  = coalesce(follow_up_note, %(note)s),
        stage           = 'outreach'
    WHERE entity_id = any(%(ids)s)
"""


def norm_url(url) -> str:
    """Comparable form of a profile URL: lowercase, no scheme / www. / trailing slash."""
    u = str(url or "").strip().lower()
    u = re.sub(r"^https?://", "", u)
    return re.sub(r"^www\.", "", u).rstrip("/")


def _merge(into: dict, info: dict) -> None:
    """Fill fields of `into` that are still empty; earlier values win."""
    for k, v in info.items():
        if v and not into.get(k):
            into[k] = v


def find_entities(cur, email: str | None, urls=()) -> list[int]:
    """entity_ids for one contacted person: each profile URL by exact source_url,
    else by the channel id / handle inside it; the email only when no URL matched."""
    ids: set[int] = set()
    for url in filter(None, urls):
        cur.execute(
            r"select entity_id from entities "
            r"where rtrim(regexp_replace(lower(source_url), '^https?://(www\.)?', ''), '/') = %s",
            (norm_url(url),),
        )
        hit = {r[0] for r in cur.fetchall()}
        keys = list(candidate_ids("", None, url) - _NOT_IDS)
        if not hit and keys:
            cur.execute(
                "select entity_id from entities "
                "where lower(platform_id) = any(%s) or lower(ltrim(handle, '@')) = any(%s)",
                (keys, keys),
            )
            hit = {r[0] for r in cur.fetchall()}
        ids |= hit
    if not ids and email:
        cur.execute("select entity_id from entities where lower(email_norm) = %s", (email,))
        ids = {r[0] for r in cur.fetchall()}
    return sorted(ids)


def mark_contacted(conn, records: list[dict], batch: str | None = None,
                   commit: bool = True) -> list[tuple[str, list[int]]]:
    """Flag the entities behind each record as contacted. A record has `to` (email),
    optional `source_url` / `urls`, `at` (datetime, default now), `batch` and `note`.
    First write wins (coalesce), so replays are no-ops -- replay history oldest-first.
    Returns (email, entity_ids) per record; an empty list means not in the DB."""
    out = []
    with conn.cursor() as cur:
        for rec in records:
            email = (rec.get("to") or "").strip().lower() or None
            ids = find_entities(cur, email, [*rec.get("urls", ()), rec.get("source_url")])
            out.append((email or "", ids))
            if not ids:
                continue
            at = rec.get("at") or datetime.now().astimezone()
            cur.execute(MARK_SQL, {
                "ids": ids, "at": at, "batch": rec.get("batch") or batch, "email": email,
                "follow": at.date() + timedelta(days=FOLLOW_UP_DAYS), "note": rec.get("note"),
            })
    if commit:
        conn.commit()
    else:
        conn.rollback()
    return out


def sheet_state(conn) -> dict[str, dict]:
    """email / profile-URL key -> {website, contacted, follow_up_at, note} for every
    entity that has a website or has been contacted."""
    state: dict[str, dict] = {}
    with conn.cursor() as cur:
        cur.execute("""
            select email_norm, source_url, coalesce(website_url, contacts->>'website'),
                   contacted_at is not null, follow_up_at, follow_up_note
            from entities
            where contacted_at is not null or coalesce(website_url, contacts->>'website') is not null
        """)
        for email, url, website, contacted, follow, note in cur.fetchall():
            info = {"website": website, "contacted": contacted, "follow_up_at": follow, "note": note}
            for key in ((email or "").lower(), norm_url(url)):
                if key:
                    _merge(state.setdefault(key, {}), info)
    return state


def yes_no_dropdown(ws, col: int) -> None:
    """Yes/No list validation on one column (rows 2..max), added once per sheet."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    letter = get_column_letter(col)
    if any(str(dv.sqref).startswith(f"{letter}2:") for dv in ws.data_validations.dataValidation):
        return
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 1000)}")


def tick_sheet(path: Path, state: dict[str, dict]) -> int:
    """Add SHEET_COLS to an .xlsx (when missing) and fill them for rows whose email
    or source_url is in `state`: "already contact" = Yes, plus website / follow up /
    note where the cell is still blank (a manual edit is never overwritten).
    Returns the number of rows ticked. PermissionError if Excel has the file open."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    header = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]

    def col(*names):
        return next((header.index(n) + 1 for n in names if n in header), None)

    c_email = col("email", "email cần gửi", "mail")
    c_src = col("source_url", "link social network", "link")
    cols = {}
    for name in SHEET_COLS:
        if name not in header:
            header.append(name)
            ws.cell(1, len(header), name)
        cols[name] = header.index(name) + 1

    ticked = 0
    for r in range(2, ws.max_row + 1):
        email = str(ws.cell(r, c_email).value or "").strip().lower() if c_email else ""
        email = email.removeprefix("mailto:").strip()
        info: dict = {}
        for key in (email, norm_url(ws.cell(r, c_src).value) if c_src else ""):
            if key in state:
                _merge(info, state[key])
        if info.get("contacted"):
            ws.cell(r, cols["already contact"]).value = "Yes"
            ticked += 1
        for name, value in (("website", info.get("website")), ("follow up", info.get("follow_up_at")),
                            ("follow up note", info.get("note"))):
            cell = ws.cell(r, cols[name])
            if value and cell.value in (None, ""):
                cell.value = value
    yes_no_dropdown(ws, cols["already contact"])
    wb.save(path)
    return ticked


def record_contacted(records: list[dict], batch: str, sheet: Path | None = None) -> None:
    """CreateMail hook: mark the DB, then tick the sheet the rows came from. Never
    raises -- the .eml files already exist; on failure it prints how to catch up."""
    follow = date.today() + timedelta(days=FOLLOW_UP_DAYS)
    state: dict[str, dict] = {}
    try:
        from db.connection import get_conn

        with get_conn() as conn:
            matches = mark_contacted(conn, records, batch)
            state = sheet_state(conn)  # DB first: an earlier contact keeps its follow-up date
        print(f"marked {sum(len(ids) for _, ids in matches)} entities contacted (batch {batch})")
        missing = [e for e, ids in matches if not ids]
        if missing:
            print("not in DB (ticked in the sheet only): " + ", ".join(missing))
    except Exception as e:  # DB down / not configured: the sheet still gets ticked
        print(f"warning: DB not updated ({e}); later run: python -m scripts.backfill_contacted")
    for rec in records:
        for key in (rec["to"].strip().lower(), norm_url(rec.get("source_url"))):
            if key:
                _merge(state.setdefault(key, {}), {"contacted": True, "follow_up_at": follow})
    if sheet:
        try:
            print(f"{sheet}: {tick_sheet(sheet, state)} rows marked already contact = Yes")
        except PermissionError:
            print(f"warning: {sheet} is open in Excel -- close it, then run: "
                  f'python -m outreach.contacted "{sheet}"')


TRACKER_COLS = ["follow up", "follow up note", "contacted_at", "batch", "platform", "name",
                "handle", "email", "phone", "website", "source_url"]


def write_tracker(conn, out: Path) -> int:
    """Everyone contacted, soonest follow-up first -> .xlsx worklist."""
    from openpyxl import Workbook

    with conn.cursor() as cur:
        cur.execute("""
            select follow_up_at, follow_up_note, contacted_at::date, contacted_batch, platform,
                   display_name, handle, email_norm, phone_norm,
                   coalesce(website_url, contacts->>'website'), source_url
            from entities where contacted_at is not null
            order by follow_up_at nulls last, contacted_at
        """)
        rows = cur.fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "follow up"
    ws.append(TRACKER_COLS)
    for row in rows:
        ws.append(list(row))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return len(rows)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="website / already contact / follow up columns from the DB")
    ap.add_argument("xlsx", nargs="*", type=Path, help="sheets to add/fill the columns in")
    ap.add_argument("--tracker", type=Path, default=None, help="also write a follow-up worklist .xlsx")
    args = ap.parse_args()
    if not args.xlsx and not args.tracker:
        ap.error("give at least one sheet or --tracker")

    from db.connection import get_conn

    with get_conn() as conn:
        state = sheet_state(conn)
        if args.tracker:
            print(f"{write_tracker(conn, args.tracker)} contacted -> {args.tracker}")
    for p in args.xlsx:
        try:
            print(f"{p}: {tick_sheet(p, state)} rows already contact = Yes")
        except PermissionError:
            print(f"{p}: open in Excel -- close it and re-run")


if __name__ == "__main__":
    main()
