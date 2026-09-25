"""Sync manually-typed contact info from a shortlist/outreach .xlsx back into the DB.

Scraped contacts flow one way: DB -> sheet (discovery.shortlist export,
outreach.contacted's tick_sheet). When you hand-find an email/phone/zalo/etc and
type it straight into the sheet, nothing writes it back -- the next scrape
refresh only ever fills what the scraper itself found. This closes that loop:
match each row to its entity by source_url (the same key outreach.contacted
uses), fill whichever DB fields are still empty, and report the rest. A field
already set in the DB to something else is left alone unless --overwrite says
the sheet should win.

    # see what would change, touch nothing:
    python -m outreach.sheet_sync results/ib_master.xlsx --dry-run
    # fill empty DB fields from the sheet:
    python -m outreach.sheet_sync results/ib_master.xlsx
    # let the sheet replace a DB value that's already set:
    python -m outreach.sheet_sync results/ib_master.xlsx --overwrite

Rows with no source_url are skipped -- there's nothing to match them to.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field as dc_field
from pathlib import Path

from discovery.common.contacts import EMAIL_RE
from outreach.contacted import find_entities

CONTACT_FIELDS = ("zalo", "telegram", "facebook", "instagram")


def _norm_email(value) -> str | None:
    s = str(value or "").strip().lower().removeprefix("mailto:").strip()
    return s if EMAIL_RE.fullmatch(s) else None


def _phone_digits(value) -> str:
    """Digits (+ leading '+') only, so '0203 655 779' and '0203655779' compare equal --
    phone_norm keeps whatever separators the scraper found in the bio text verbatim."""
    return re.sub(r"[^\d+]", "", str(value or ""))


def _norm_phone(value) -> str | None:
    if value in (None, ""):
        return None
    s = str(int(value)) if isinstance(value, (int, float)) else str(value).strip()
    digits = re.sub(r"[^\d+]", "", s)
    if not digits:
        return None
    if not digits.startswith(("+", "0")) and len(digits) == 9:
        digits = "0" + digits  # Excel drops the leading 0 on numeric-typed phone cells
    return digits


def _read_rows(path: Path) -> list[dict]:
    """Non-empty cells per row that has a source_url, keyed by email/phone/website/
    the CONTACT_FIELDS -- only columns actually present in the sheet are read."""
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True).active  # not read_only: that mode keeps the file locked on Windows
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip().lower() for c in (next(rows, []) or [])]

    def col(*names):
        return next((header.index(n) for n in names if n in header), None)

    idx = {"source_url": col("source_url", "link social network", "link"),
           "email": col("email", "email cần gửi", "mail"),
           "phone": col("phone",), "website": col("website",),
           **{k: col(k) for k in CONTACT_FIELDS}}

    out = []
    for row in rows:
        rec = {}
        for key, i in idx.items():
            if i is not None and i < len(row) and row[i] not in (None, ""):
                rec[key] = row[i]
        if rec.get("source_url"):
            out.append(rec)
    return out


@dataclass
class SyncReport:
    filled: list = dc_field(default_factory=list)
    overwritten: list = dc_field(default_factory=list)
    conflicts: list = dc_field(default_factory=list)
    invalid: list = dc_field(default_factory=list)
    unmatched: list = dc_field(default_factory=list)

    def print(self) -> None:
        print(f"  filled {len(self.filled)}, overwrote {len(self.overwritten)}, "
              f"{len(self.conflicts)} conflicts, {len(self.invalid)} invalid emails, "
              f"{len(self.unmatched)} rows not in DB")
        for url, key, old, new in self.conflicts:
            print(f"  conflict {url} {key}: db={old!r} sheet={new!r} (kept db; rerun --overwrite to replace)")
        for url, value in self.invalid:
            print(f"  invalid email {url}: {value!r}")
        for url in self.unmatched:
            print(f"  not in DB: {url}")


def sync_sheet(conn, path: Path, overwrite: bool = False, dry_run: bool = False) -> SyncReport:
    from psycopg.types.json import Jsonb

    report = SyncReport()
    with conn.cursor() as cur:
        for row in _read_rows(path):
            url = row["source_url"]
            ids = find_entities(cur, None, [url])
            if not ids:
                report.unmatched.append(url)
                continue

            email = _norm_email(row["email"]) if "email" in row else None
            if "email" in row and email is None:
                report.invalid.append((url, row["email"]))
            sheet_values = {
                "email": email,
                "phone": _norm_phone(row.get("phone")) if "phone" in row else None,
                "website": str(row["website"]).strip() if "website" in row else None,
                **{k: str(row[k]).strip() for k in CONTACT_FIELDS if k in row},
            }

            for entity_id in ids:
                cur.execute("select email_norm, phone_norm, website_url, contacts "
                            "from entities where entity_id = %s", (entity_id,))
                cur_email, cur_phone, cur_website, contacts = cur.fetchone()
                current = {"email": cur_email, "phone": cur_phone, "website": cur_website,
                           **(contacts or {})}
                final_email, final_phone, final_website = cur_email, cur_phone, cur_website
                new_contacts = dict(contacts or {})
                changed = False

                for key, new in sheet_values.items():
                    if not new:
                        continue
                    old = current.get(key)
                    same = _phone_digits(old) == new if key == "phone" else old == new
                    if same:
                        continue
                    if old and not overwrite:
                        report.conflicts.append((url, key, old, new))
                        continue
                    (report.overwritten if old else report.filled).append((url, key, old, new))
                    changed = True
                    if key == "email":
                        final_email = new
                    elif key == "phone":
                        final_phone = new
                    elif key == "website":
                        final_website = new
                    else:
                        new_contacts[key] = new

                if dry_run or not changed:
                    continue
                cur.execute(
                    "update entities set email_norm=%s, phone_norm=%s, website_url=%s, "
                    "contacts=%s where entity_id=%s",
                    (final_email, final_phone, final_website,
                     Jsonb(new_contacts) if new_contacts else None, entity_id),
                )
    if dry_run:
        conn.rollback()
    else:
        conn.commit()
    return report


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Sync manually-entered sheet contacts back into the DB")
    ap.add_argument("xlsx", nargs="+", type=Path)
    ap.add_argument("--overwrite", action="store_true", help="let the sheet replace a differing DB value")
    ap.add_argument("--dry-run", action="store_true", help="report changes without writing")
    args = ap.parse_args()

    from db.connection import get_conn

    with get_conn() as conn:
        for path in args.xlsx:
            print(f"{path}:")
            sync_sheet(conn, path, overwrite=args.overwrite, dry_run=args.dry_run).print()


if __name__ == "__main__":
    main()
