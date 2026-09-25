"""Backfill "already contacted" for outreach sent before CreateMail recorded it.

Replays two sources through outreach.contacted.mark_contacted, oldest first so the
earliest contact date wins (re-running is a no-op):
  * every outreach/eml/<label>/*.eml -- To + Date header, batch = folder name. An
    .eml carries only the address, so its profile URL is looked up in the local
    sheets that have email + source_url (results/**, outreach/, ~/Downloads).
  * the Outlook Sent Items export (default ~/Downloads/TMGM_Sent_enriched.xlsx:
    Time / mail / facebook / tiktok / youtube / Client reply), batch 'sent-items'.

    # show what matches which entity, write nothing:
    python -m scripts.backfill_contacted --dry-run
    python -m scripts.backfill_contacted
"""

import argparse
import sys
from datetime import datetime
from email import message_from_binary_file, policy
from email.utils import parsedate_to_datetime
from pathlib import Path

from outreach.contacted import mark_contacted
from outreach.eml import EML_ROOT

DOWNLOADS = Path.home() / "Downloads"
SHEET_DIRS = [Path("results"), Path("outreach"), DOWNLOADS]


def _rows(xlsx: Path):
    """(lowercased header, row tuples) of a sheet's first tab; ([], []) if unreadable."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
    except Exception:
        return [], []
    try:  # read_only keeps the file locked on Windows until closed
        rows = wb.active.iter_rows(values_only=True)
        return [str(c or "").strip().lower() for c in next(rows, ())], list(rows)
    finally:
        wb.close()


def email_url_index(dirs=SHEET_DIRS) -> dict[str, set[str]]:
    """email -> source_urls, from every sheet under `dirs` with both columns."""
    index: dict[str, set[str]] = {}
    for d in dirs:
        for f in d.rglob("*.xlsx") if d != DOWNLOADS else d.glob("*.xlsx"):
            if f.name.startswith("~$"):
                continue
            header, rows = _rows(f)
            if "email" not in header or "source_url" not in header:
                continue
            e, s = header.index("email"), header.index("source_url")
            for row in rows:
                email = str(row[e] or "").strip().lower().removeprefix("mailto:").strip()
                if email and row[s]:
                    index.setdefault(email, set()).add(str(row[s]).strip())
    return index


def eml_records(root: Path, index: dict[str, set[str]]) -> list[dict]:
    out = []
    for f in sorted(root.glob("*/*.eml")):
        with f.open("rb") as fh:
            msg = message_from_binary_file(fh, policy=policy.default)
        to = str(msg["To"] or "").strip().lower()
        at = parsedate_to_datetime(str(msg["Date"])) if msg["Date"] else \
            datetime.fromtimestamp(f.stat().st_mtime).astimezone()
        out.append({"to": to, "urls": sorted(index.get(to, ())), "at": at, "batch": f.parent.name})
    return out


def sent_records(xlsx: Path) -> list[dict]:
    header, rows = _rows(xlsx)
    if not header:
        return []
    ix = {name: header.index(name) for name in header}
    out = []
    for row in rows:
        get = lambda name: str(row[ix[name]] or "").strip() if name in ix else ""
        to = get("mail").strip("'").lower()
        if not to or not get("time"):
            continue
        out.append({
            "to": to,
            "urls": [u for u in (get("facebook"), get("tiktok"), get("youtube")) if u],
            "at": datetime.strptime(get("time"), "%Y-%m-%d %H:%M").astimezone(),
            "batch": "sent-items",
            "note": "replied" if get("client reply").lower() == "yes" else None,
        })
    return out


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Backfill already-contacted from old .eml batches + Sent Items")
    ap.add_argument("--eml-root", type=Path, default=EML_ROOT)
    ap.add_argument("--sent", type=Path, default=DOWNLOADS / "TMGM_Sent_enriched.xlsx",
                    help="Sent Items export (Time/mail/facebook/tiktok/youtube/Client reply)")
    ap.add_argument("--dry-run", action="store_true", help="show matches, write nothing")
    args = ap.parse_args()

    records = eml_records(args.eml_root, email_url_index())
    if args.sent.exists():
        records += sent_records(args.sent)
    else:
        print(f"warning: {args.sent} not found, only .eml batches")
    records.sort(key=lambda r: r["at"])

    from db.connection import get_conn

    with get_conn() as conn:
        matches = mark_contacted(conn, records, commit=not args.dry_run)
    missing = []
    for rec, (email, ids) in zip(records, matches):
        if ids:
            print(f"{rec['at']:%Y-%m-%d}  {rec['batch']:<28} {email:<40} -> {ids}")
        else:
            missing.append(f"{email} ({rec['batch']})")
    matched = {i for _, ids in matches for i in ids}
    print(f"\n{len(records)} records -> {len(matched)} entities"
          f"{' (dry run, nothing written)' if args.dry_run else ' marked contacted'}")
    if missing:
        print(f"{len(missing)} not in DB: " + ", ".join(missing))


if __name__ == "__main__":
    main()
